#!/usr/bin/env python3
"""Attenda Team Engagement Offer Letters — branded package (5 documents).
Content source: user-approved draft (September 2026 premium model)."""
import datetime, os, subprocess, glob
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BRAND = "/tmp/attenda_brand"
SVG_DIR = os.path.expanduser("~/Projects/attenda/public/brand")
TEAL = RGBColor(0x15, 0x8A, 0x7C)
DARK = RGBColor(0x1A, 0x2A, 0x3A)
GRAY = RGBColor(0x6B, 0x72, 0x80)
MINT_FILL = "E8F4F1"
PALE_FILL = "F4FAF8"
TODAY = datetime.date(2026, 9, 2)
START = datetime.date(2026, 9, 7)
SIGNBY = datetime.date(2026, 9, 4)
fmt = lambda d: d.strftime("%B %d, %Y").replace(" 0", " ")

# ---------------- assets (self-healing) ----------------
def ensure_assets():
    os.makedirs(BRAND, exist_ok=True)
    logo = f"{BRAND}/logo_primary_cropped.png"
    if not os.path.exists(logo):
        subprocess.run(["qlmanage", "-t", "-s", "2000", "-o", BRAND,
                        f"{SVG_DIR}/logo-primary.svg"], capture_output=True)
        src = glob.glob(f"{BRAND}/logo-primary.svg.png")[0]
        from PIL import Image
        im = Image.open(src).convert("RGBA")
        bbox = im.getbbox()
        im.crop(bbox).save(logo)
    rule = f"{BRAND}/gradient_rule.png"
    if not os.path.exists(rule):
        from PIL import Image
        stops = [(0x5E, 0xCF, 0xC0), (0x1E, 0x9E, 0x8F), (0x15, 0x8A, 0x7C), (0x0E, 0x6B, 0x60)]
        W, H = 2400, 56
        im = Image.new("RGB", (W, H))
        px = im.load()
        for x in range(W):
            t = x / (W - 1) * (len(stops) - 1)
            i = min(int(t), len(stops) - 2)
            f = t - i
            c = tuple(int(a + (b - a) * f) for a, b in zip(stops[i], stops[i + 1]))
            for y in range(H):
                px[x, y] = c
        im.save(rule)
    return f"{BRAND}/logo_primary_cropped.png", rule

LOGO, RULE = ensure_assets()

# ---------------- doc helpers ----------------
def cell_shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:fill"), fill)
    tcPr.append(shd)

def cell_margins(cell, top=100, bottom=100, left=160, right=160):
    tcPr = cell._tc.get_or_add_tcPr()
    m = OxmlElement("w:tcMar")
    for tag, val in (("top", top), ("bottom", bottom), ("start", left), ("end", right)):
        el = OxmlElement(f"w:{tag}"); el.set(qn("w:w"), str(val)); el.set(qn("w:type"), "dxa")
        m.append(el)
    tcPr.append(m)

def no_table_borders(tbl):
    tblPr = tbl._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}"); el.set(qn("w:val"), "none")
        borders.append(el)
    tblPr.append(borders)

def spacer(doc, pts=10):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(pts)
    pPr = p._p.get_or_add_pPr()
    rPr = OxmlElement("w:rPr"); sz = OxmlElement("w:sz"); sz.set(qn("w:val"), "12")
    rPr.append(sz); pPr.append(rPr)

def heading(doc, text, before=12, after=5):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before); p.paragraph_format.space_after = Pt(after)
    r = p.add_run(text)
    r.font.name = "Georgia"; r.font.size = Pt(13.5); r.font.bold = True; r.font.color.rgb = TEAL
    return p

def boxed_callout(doc, runs, sub=None, italic=None, fill=MINT_FILL):
    tbl = doc.add_table(rows=1, cols=1)
    no_table_borders(tbl)
    cell = tbl.rows[0].cells[0]
    cell.width = Inches(7.0)
    cell_shade(cell, fill)
    cell_margins(cell, top=150, bottom=150, left=220, right=220)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    cp = cell.paragraphs[0]
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.space_after = Pt(2)
    for text, size, bold, color in runs:
        r = cp.add_run(text)
        r.font.name = "Georgia"; r.font.size = Pt(size); r.font.bold = True; r.font.color.rgb = color
    if sub:
        cp2 = cell.add_paragraph(); cp2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp2.paragraph_format.space_after = Pt(3)
        r = cp2.add_run(sub); r.font.size = Pt(10.5); r.font.color.rgb = GRAY
    if italic:
        cp3 = cell.add_paragraph(); cp3.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = cp2b = cp2 = cell.paragraphs[-1]
        cp3 = cell.add_paragraph(); cp3.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = cp3.add_run(italic); r.font.size = Pt(10.5); r.font.italic = True; r.font.color.rgb = GRAY
    return tbl

def fact_strip(doc, facts, cols=None):
    spacer(doc, 8)
    cols = cols or len(facts)
    tbl = doc.add_table(rows=2, cols=cols)
    no_table_borders(tbl)
    width = Inches(7.0 / cols)
    for i, (label, value) in enumerate(facts):
        c, c2 = tbl.rows[0].cells[i], tbl.rows[1].cells[i]
        for cc in (c, c2):
            cc.width = width; cell_shade(cc, PALE_FILL)
            cell_margins(cc, top=70, bottom=70, left=120, right=120)
        rl = c.paragraphs[0].add_run(label)
        rl.font.size = Pt(8.5); rl.font.bold = True; rl.font.color.rgb = GRAY
        rv = c2.paragraphs[0].add_run(value)
        rv.font.size = Pt(10.5); rv.font.bold = True; rv.font.color.rgb = DARK

def build_letter(cfg):
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"; st.font.size = Pt(10.5); st.font.color.rgb = DARK
    st.paragraph_format.space_after = Pt(7)
    sec = doc.sections[0]
    sec.top_margin = Inches(0.55); sec.bottom_margin = Inches(0.55)
    sec.left_margin = Inches(0.75); sec.right_margin = Inches(0.75)

    g = doc.add_paragraph(); g.paragraph_format.space_after = Pt(12)
    g.add_run().add_picture(RULE, width=Inches(7.0), height=Inches(0.09))

    lp = doc.add_paragraph(); lp.paragraph_format.space_after = Pt(2)
    lp.add_run().add_picture(LOGO, width=Inches(2.35))

    kick = doc.add_paragraph(); kick.paragraph_format.space_after = Pt(10)
    rk = kick.add_run(cfg["kicker"])
    rk.font.name = "Georgia"; rk.font.size = Pt(9.5); rk.font.bold = True; rk.font.color.rgb = GRAY
    r2 = kick.add_run("     ·     CONFIDENTIAL")
    r2.font.name = "Georgia"; r2.font.size = Pt(9.5); r2.font.color.rgb = GRAY
    pPr = kick._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr"); b = OxmlElement("w:bottom")
    b.set(qn("w:val"), "single"); b.set(qn("w:sz"), "6"); b.set(qn("w:space"), "6"); b.set(qn("w:color"), "CBE4DF")
    pBdr.append(b); pPr.append(pBdr)

    d = doc.add_paragraph(fmt(TODAY)); d.paragraph_format.space_after = Pt(12)

    if cfg.get("greeting"):
        h = doc.add_paragraph()
        rh = h.add_run(cfg["greeting"])
        rh.font.name = "Georgia"; rh.font.size = Pt(13); rh.font.bold = True
        h.paragraph_format.space_after = Pt(8)

    w = doc.add_paragraph()
    parts = cfg["intro"]
    for text, bold in parts:
        r = w.add_run(text); r.font.bold = bold

    # comp / role callout
    if cfg.get("comp_runs"):
        boxed_callout(doc, cfg["comp_runs"], sub=cfg.get("comp_sub"), italic=cfg.get("comp_italic"))
        spacer(doc, 4)

    if cfg.get("facts"):
        fact_strip(doc, cfg["facts"])

    heading(doc, cfg["duties_heading"], before=12, after=5)
    for title, body in cfg["duties"]:
        bp = doc.add_paragraph(style="List Bullet")
        bp.paragraph_format.space_after = Pt(5)
        if title:
            rt = bp.add_run(title + " "); rt.font.bold = True
        bp.add_run(body)

    heading(doc, cfg["boundaries_heading"])
    bp = doc.add_paragraph(cfg["boundaries"])
    bp.paragraph_format.space_after = Pt(6)

    if cfg.get("growth"):
        gp = doc.add_paragraph(cfg["growth"])
        gp.paragraph_format.space_after = Pt(6)

    if cfg.get("recognition"):
        boxed_callout(doc, cfg["recognition"][0], sub=None, italic=cfg["recognition"][1], fill=PALE_FILL)

    doc.add_paragraph(cfg["closing"]).paragraph_format.space_after = Pt(6)

    doc.add_paragraph("Sincerely,").paragraph_format.space_after = Pt(14)
    sig = doc.add_paragraph()
    rs = sig.add_run("Alejandro Soria"); rs.font.bold = True; rs.font.name = "Georgia"; rs.font.size = Pt(13)
    doc.add_paragraph("Founder, Attenda").paragraph_format.space_after = Pt(12)

    if cfg.get("ack", True):
        spacer(doc, 6)
        tbl = doc.add_table(rows=2, cols=1)
        no_table_borders(tbl)
        hc, bc = tbl.rows[0].cells[0], tbl.rows[1].cells[0]
        for cc in (hc, bc):
            cc.width = Inches(7.0)
        cell_shade(hc, MINT_FILL); cell_margins(hc, top=80, bottom=80, left=200, right=200)
        rr = hc.paragraphs[0].add_run("ACKNOWLEDGMENT")
        rr.font.size = Pt(9); rr.font.bold = True; rr.font.color.rgb = TEAL
        cell_margins(bc, top=100, bottom=100, left=200, right=200)
        p1 = bc.paragraphs[0]
        r = p1.add_run("I acknowledge that I have reviewed the initial scope and compensation described above. "
                       "Final working arrangements, payment terms, confidentiality obligations, and any applicable "
                       "contractor or employment terms will be documented separately as required.")
        r.font.size = Pt(9.5); r.font.color.rgb = GRAY
        p1.paragraph_format.space_after = Pt(10)
        p2 = bc.add_paragraph()
        r = p2.add_run("Name: ____________________________        Signature: ____________________________        Date: ______________")
        r.font.size = Pt(10.5)

    fp = doc.add_paragraph(); fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.paragraph_format.space_before = Pt(12)
    fp.add_run().add_picture(RULE, width=Inches(7.0), height=Inches(0.09))
    fp2 = doc.add_paragraph(); fp2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rf = fp2.add_run("attenda  ·  premium hotel operations  ·  attendaapp.com")
    rf.font.size = Pt(9); rf.font.color.rgb = GRAY

    doc.save(cfg["out"])
    print("saved:", cfg["out"])

# ---------------- configurations (verbatim draft copy) ----------------
ACK_NOTE = "I acknowledge that I have reviewed the initial scope and compensation described above."

PEOPLE = []

PEOPLE.append({
    "out": "/Users/thrilzco/Documents/Attenda_Offer_Juan_Field_Operations_Manager.docx",
    "kicker": "TEAM ENGAGEMENT OFFER — FIELD OPERATIONS",
    "greeting": "Dear Juan,",
    "intro": [("I am pleased to invite you to support Attenda as ", False),
              ("Field Operations Manager", True),
              (" as we launch our premium hotel operations service. You will be Attenda's boots on the ground, "
               "helping turn operating priorities into consistent property-level execution.", False)],
    "comp_runs": [("$1,000", 34, True, TEAL), ("  per month", 13, True, DARK)],
    "comp_sub": "For an expected commitment of approximately 10 hours per week · paid biweekly",
    "comp_italic": "As the portfolio grows, the responsibilities and compensation associated with this function may be reviewed.",
    "facts": [("START DATE", fmt(START)), ("SIGN BY", fmt(SIGNBY)), ("PAY CYCLE", "Biweekly"), ("REPORTS TO", "Controller & Founder")],
    "duties_heading": "Initial responsibilities",
    "duties": [
        (None, "Maintain a regular on-property presence, generally totaling approximately 10 hours per week."),
        (None, "Follow up on projects, assignments, inspections, standards, and operational priorities established by Attenda."),
        (None, "Verify completion and quality rather than relying only on reported completion."),
        (None, "Identify operational gaps, risks, recurring issues, and opportunities for improvement."),
        (None, "Support the hotel manager with implementation, organization, and accountability."),
        (None, "Provide concise field updates and escalate significant matters to the Controller and Alejandro."),
        (None, "Use Attenda as the primary workflow and accountability system for assigned work."),
    ],
    "boundaries_heading": "Role boundaries",
    "boundaries": ("The hotel's manager remains responsible for the hotel's employees and day-to-day property operation. "
                   "Your role is not to replace the hotel manager or independently direct the client's employees outside the agreed "
                   "operating structure. Your function is field support, verification, execution follow-up, and accountability on behalf of Attenda."),
    "growth": ("This first property will help establish the field-operations model that Attenda can replicate across additional hotels."),
    "closing": "I am excited to have you help build the operating foundation of Attenda.",
})

PEOPLE.append({
    "out": "/Users/thrilzco/Documents/Attenda_Offer_Drashti_Fractional_Controller.docx",
    "kicker": "OFFER OF ENGAGEMENT — FINANCE",
    "greeting": "Dear Drashti,",
    "intro": [("I am pleased to invite you to support Attenda in a ", False),
              ("fractional controller and financial-oversight", True),
              (" capacity. Your role is to give leadership a clear financial picture so operating decisions are "
               "supported by disciplined financial review.", False)],
    "comp_runs": [("$600", 34, True, TEAL), ("  per month", 13, True, DARK)],
    "comp_sub": "For the initial agreed fractional scope",
    "comp_italic": "As Attenda adds properties and the financial function expands, we can revisit scope, workload, title, and compensation.",
    "facts": [("START DATE", fmt(START)), ("SIGN BY", fmt(SIGNBY)), ("REPORTS TO", "Founder, Attenda")],
    "duties_heading": "Initial responsibilities",
    "duties": [
        (None, "Review agreed hotel financial reports, operating results, expenses, and material variances."),
        (None, "Identify financial risks, unusual trends, control concerns, and potential savings opportunities."),
        (None, "Provide concise financial observations and recommendations to Alejandro."),
        (None, "Help connect operating activity to financial performance and highlight areas requiring management attention."),
        (None, "Support agreed reconciliations, controls, or financial reviews within the defined scope."),
        (None, "Help establish repeatable financial-review processes that can scale to additional Attenda clients."),
    ],
    "boundaries_heading": "Role boundaries",
    "boundaries": ("This is an oversight and advisory function within the agreed scope. You are not responsible for running the hotel's "
                   "daily operation, and the engagement does not automatically transfer the hotel's accounting, payroll, tax, audit, or "
                   "fiduciary responsibilities to you or Attenda."),
    "closing": "I look forward to building a disciplined financial foundation with you.",
})

PEOPLE.append({
    "out": "/Users/thrilzco/Documents/Attenda_Offer_Su_Sales_Business_Development.docx",
    "kicker": "OFFER OF ENGAGEMENT — SALES & BUSINESS DEVELOPMENT",
    "greeting": "Dear Su,",
    "intro": [("I am pleased to invite you to support Attenda in ", False),
              ("Sales & Business Development", True),
              (" as we begin serving our first premium hotel client. Your focus will be creating measurable new business "
               "opportunities while helping us develop a repeatable hotel-sales function.", False)],
    "comp_runs": [("$600", 34, True, TEAL), ("  per month", 13, True, DARK)],
    "comp_sub": "For the initial scope, plus performance-based incentives when separately defined and approved",
    "comp_italic": ("Any commission or incentive plan will be documented separately and should define qualifying revenue, attribution, "
                    "payment timing, exclusions, and the commissionable period. No commission should be assumed unless agreed in writing."),
    "facts": [("START DATE", fmt(START)), ("SIGN BY", fmt(SIGNBY)), ("REPORTS TO", "Founder, Attenda")],
    "duties_heading": "Initial responsibilities",
    "duties": [
        (None, "Prospect for qualified local, corporate, group, partnership, and other agreed hotel business opportunities."),
        (None, "Develop and maintain an organized pipeline of prospects, contacts, opportunities, and follow-up activity."),
        (None, "Coordinate sales priorities and target accounts with Alejandro and the property team."),
        (None, "Document attributable new business so performance incentives can be measured fairly."),
        (None, "Support Attenda business development when specifically assigned, while keeping hotel-client opportunities properly separated."),
        (None, "Provide regular updates on pipeline activity, wins, losses, next steps, and obstacles."),
    ],
    "boundaries_heading": "Role boundaries",
    "boundaries": ("No commission should be assumed unless it has been agreed in writing. Any incentive plan will define qualifying revenue, "
                   "attribution, payment timing, exclusions, and the period during which an account or booking remains commissionable."),
    "growth": ("As Attenda expands, this function can grow into a broader revenue and business-development department with increased responsibility and upside."),
    "closing": "I am excited to have you help build the revenue engine behind the model.",
})

PEOPLE.append({
    "out": "/Users/thrilzco/Documents/Attenda_Offer_Francisca_Onboarding_Training.docx",
    "kicker": "OFFER OF ENGAGEMENT — ONBOARDING & TRAINING",
    "greeting": "Dear Francisca,",
    "intro": [("I am pleased to invite you to support Attenda in ", False),
              ("Onboarding, Training & Client Relations", True),
              (". Your immediate priority is to ensure that our first client's property leadership understands how to operate "
               "successfully within the Attenda system.", False)],
    "comp_runs": [("To be confirmed", 26, True, TEAL)],
    "comp_sub": "Initial compensation and/or project-based payment — to be confirmed based on the agreed training and client-support scope",
    "comp_italic": ("The training process created at this first property should become the foundation for onboarding and certifying "
                    "leaders across future Attenda properties."),
    "facts": [("START DATE", fmt(START)), ("SIGN BY", fmt(SIGNBY)), ("REPORTS TO", "Founder, Attenda")],
    "duties_heading": "Initial responsibilities",
    "duties": [
        (None, "Lead initial Attenda onboarding for designated property leaders and users."),
        (None, "Train users on platform workflows, assigned SOPs, checklists, inspections, reporting, escalation, and accountability standards."),
        (None, "Coordinate the initial training week and document completion of required learning activities."),
        (None, "Assess practical understanding before recommending certification."),
        (None, "Support adoption after launch and identify areas requiring retraining or clarification."),
        (None, "Maintain organized training materials and help build a repeatable Attenda certification program."),
        (None, "Serve as a client-relations resource for training, adoption, and appropriate support matters."),
    ],
    "boundaries_heading": "Role boundaries",
    "boundaries": ("Training and certification authorize a client's team to operate within Attenda's system and standards; they do not make "
                   "the client's employees employees of Attenda. Employment supervision, compensation, disciplinary decisions, and other "
                   "employer responsibilities remain with the hotel."),
    "closing": "I am excited to have you build the training and adoption standard that will support Attenda's growth.",
})

PEOPLE.append({
    "out": "/Users/thrilzco/Documents/Attenda_Certification_Daquan_Property_Leader.docx",
    "kicker": "PROPERTY LEADER CERTIFICATION FRAMEWORK",
    "greeting": None,
    "intro": [("Daquan", True),
              (" remains employed directly by the Best Western property and is not being offered employment by Attenda. Because she is the "
               "day-to-day hotel leader operating within Attenda's premium service model, she will complete Attenda property-leader training.", False)],
    "comp_runs": [("Certification Framework", 24, True, TEAL)],
    "comp_sub": "Attenda training in the operating model — employment, compensation, and employer responsibilities remain with the hotel",
    "facts": [("PROPERTY", "Best Western — client property"), ("RELATIONSHIP", "Employed by the hotel"), ("RECOGNITION", "Attenda Certified Property Leader")],
    "duties_heading": "Initial certification expectations",
    "duties": [
        (None, "Complete the initial Attenda training program led by Francisca."),
        (None, "Demonstrate proficiency with the platform, assigned workflows, SOPs, checklists, inspections, reporting, and escalation procedures."),
        (None, "Understand the operating flow between hotel staff, property leadership, Attenda field operations, and executive oversight."),
        (None, "Complete practical exercises or assessments established for the property-leader role."),
        (None, "Maintain certification through continued proper use of Attenda and any required refresher training."),
    ],
    "boundaries_heading": "What certification means",
    "boundaries": ("Successful completion may be recognized internally as Attenda Certified Property Leader for the participating property. "
                   "Certification reflects training in the Attenda operating model only. It does not alter Daquan's employer, create an employment "
                   "relationship with Attenda, or transfer the hotel's employer responsibilities to Attenda."),
    "ack": False,
    "closing": "Welcome to the Attenda operating model — we are glad to have you leading the first property through it.",
})

for cfg in PEOPLE:
    build_letter(cfg)

print("package complete:", len(PEOPLE), "documents")