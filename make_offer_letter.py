#!/usr/bin/env python3
"""Generate Attenda offer letter (Word) for Drashti — Executive Operations Assistant, $600/hotel/mo."""
import datetime
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

TEAL = RGBColor(0x15, 0x8A, 0x7C)          # #158A7C brand teal
TEAL_HEX = "158A7C"
DARK = RGBColor(0x1A, 0x2A, 0x3A)          # near-black text
GRAY = RGBColor(0x6B, 0x72, 0x80)          # secondary gray

CAND_FULL = "Drashti"
TODAY = datetime.date(2026, 9, 2)
ACCEPT_BY = datetime.date(2026, 9, 4)       # Friday
START = datetime.date(2026, 9, 7)           # next Monday


def fmt_date(d):
    return d.strftime("%B %d, %Y").replace(" 0", " ")


def bottom_border(paragraph, color=TEAL_HEX, sz="16"):
    p = paragraph._p
    pPr = p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), sz)
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), color)
    pBdr.append(bottom)
    pPr.append(pBdr)


def shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


doc = Document()

# Base style
style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(11)
style.font.color.rgb = DARK
style.paragraph_format.space_after = Pt(8)

for section in doc.sections:
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)

# ---------- LETTERHEAD ----------
lh = doc.add_paragraph()
lh.alignment = WD_ALIGN_PARAGRAPH.LEFT
r = lh.add_run("attenda")
r.font.name = "Calibri"
r.font.size = Pt(30)
r.font.bold = True
r.font.color.rgb = TEAL

tag = doc.add_paragraph()
tag.paragraph_format.space_after = Pt(2)
rt = tag.add_run("Hotel Operations Platform for Independent Properties")
rt.font.size = Pt(11)
rt.font.color.rgb = GRAY

contact = doc.add_paragraph()
contact.paragraph_format.space_after = Pt(6)
rc = contact.add_run("attendaapp.com  ·  support@attendaapp.com  ·  Miami, FL")
rc.font.size = Pt(9.5)
rc.font.color.rgb = GRAY
bottom_border(contact)

doc.add_paragraph()

# ---------- DATE + RECIPIENT ----------
p = doc.add_paragraph(fmt_date(TODAY))
p.paragraph_format.space_after = Pt(14)

p = doc.add_paragraph(CAND_FULL)
p.paragraph_format.space_after = Pt(0)

# ---------- OPENING ----------
p = doc.add_paragraph("Dear " + CAND_FULL + ",")
p.paragraph_format.space_after = Pt(10)

doc.add_paragraph(
    "We are pleased to offer you the position of Executive Operations Assistant with Attenda, working "
    "hand-in-hand with the Founder to keep every property we serve organized, informed, and ahead of problems. "
    "This is a hands-on role at the center of the operation — not a back-office position — and it is a founding "
    "seat in a company that is actively growing."
)

# ---------- ROLE & RESPONSIBILITIES ----------
h = doc.add_paragraph()
h.paragraph_format.space_before = Pt(6)
rh = h.add_run("Role & Responsibilities")
rh.font.size = Pt(13)
rh.font.bold = True
rh.font.color.rgb = TEAL

duties = [
    "Prepare and distribute the daily operations report for every active property — occupancy, arrivals and departures, incidents, work orders, and anything needing follow-up — so nothing slips through the cracks.",
    "Review staffing schedules against forecast and actual demand; flag labor overruns or gaps early and bring concrete suggestions before they hit the P&L.",
    "Own the monthly maintenance calendar for each property — schedule it, track it, chase completion, and keep preventive items ahead of breakdowns.",
    "Build the monthly owner presentation for each property — performance, wins, issues, and the plan for next month — in a clean format owners actually read.",
    "Think ahead on the Founder's behalf: track loose ends, chase follow-ups, and keep checklists and open items moving across properties.",
    "Support day-to-day coordination between properties, staff, and vendors — hands-on, in the weeds, where execution matters more than theory.",
]
for d in duties:
    bp = doc.add_paragraph(d, style="List Bullet")
    bp.paragraph_format.space_after = Pt(4)

note = doc.add_paragraph(
    "In short: you are the person who helps the Founder keep every property in check — reporting, labor, "
    "maintenance, and owners — with an eye always on what's coming next, not just what happened yesterday."
)
note.paragraph_format.space_before = Pt(4)

# ---------- COMPENSATION ----------
h = doc.add_paragraph()
h.paragraph_format.space_before = Pt(6)
rh = h.add_run("Compensation")
rh.font.size = Pt(13)
rh.font.bold = True
rh.font.color.rgb = TEAL

comp = doc.add_paragraph("You will be paid a monthly retainer of ")
rc = comp.add_run("$600.00 per hotel, per month")
rc.font.bold = True
comp.add_run(", for each active property you support. Payment is made monthly.")

growth = doc.add_paragraph(
    "Attenda is an early-stage company onboarding its first properties, so this role is intentionally built "
    "per property instead of as a fixed salary: each new hotel added to the portfolio increases your monthly "
    "retainer by $600.00. As Attenda grows, this engagement grows with it — and we will revisit compensation "
    "and structure together in good faith as the company scales."
)

# ---------- START DATE ----------
h = doc.add_paragraph()
h.paragraph_format.space_before = Pt(6)
rh = h.add_run("Start Date")
rh.font.size = Pt(13)
rh.font.bold = True
rh.font.color.rgb = TEAL

doc.add_paragraph(
    f"Upon acceptance of this offer, your start date will be Monday, {fmt_date(START)}. "
    f"Please sign and return this letter no later than {fmt_date(ACCEPT_BY)} so we can hit the ground running."
)

# ---------- CONFIDENTIALITY ----------
h = doc.add_paragraph()
h.paragraph_format.space_before = Pt(6)
rh = h.add_run("Confidentiality")
rh.font.size = Pt(13)
rh.font.bold = True
rh.font.color.rgb = TEAL

doc.add_paragraph(
    "You will have access to sensitive operational, staffing, and financial information belonging to Attenda "
    "and its client properties. We expect that information to remain strictly confidential, during and after "
    "your engagement."
)

# ---------- CLOSING ----------
doc.add_paragraph(
    "We are excited to have you at the center of operations as Attenda grows from one property into many. "
    "If you have any questions before signing, reach out any time."
)
doc.add_paragraph("Warm regards,")
sig = doc.add_paragraph()
sig.paragraph_format.space_after = Pt(2)
rs = sig.add_run("Alejandro Soria")
rs.font.bold = True
doc.add_paragraph("Founder, Attenda").paragraph_format.space_after = Pt(14)

# ---------- SIGNATURE BLOCKS ----------
tbl = doc.add_table(rows=2, cols=2)
tbl.autofit = False
widths = (Inches(3.4), Inches(3.4))
for row in tbl.rows:
    for cell, w in zip(row.cells, widths):
        cell.width = w
        cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP

h1, h2 = tbl.rows[0].cells
for cell, label in ((h1, "For Attenda"), (h2, "Acceptance")):
    shade(cell, "E8F2F0")
    cp = cell.paragraphs[0]
    cp.paragraph_format.space_after = Pt(2)
    rr = cp.add_run(label)
    rr.font.bold = True
    rr.font.size = Pt(10)
    rr.font.color.rgb = TEAL

b1, b2 = tbl.rows[1].cells

def sigcell(cell, lines):
    first = True
    for text in lines:
        p = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        p.paragraph_format.space_after = Pt(10)
        rr = p.add_run(text)
        rr.font.size = Pt(10.5)

sigcell(b1, ["Signature: ______________________", "", "Alejandro Soria", "Founder, Attenda", f"Date: {fmt_date(TODAY)}"])
sigcell(b2, [f"I, {CAND_FULL}, accept this offer.", "",
             "Signature: ______________________", "", f"Date: ________________"])

# ---------- FOOTER ----------
fp = doc.add_paragraph()
fp.paragraph_format.space_before = Pt(16)
bottom_border(fp, color="D0D5DD", sz="6")
fp2 = doc.add_paragraph()
fp2.alignment = WD_ALIGN_PARAGRAPH.CENTER
rf = fp2.add_run("attenda — the operations platform for independent hotels")
rf.font.size = Pt(9)
rf.font.color.rgb = GRAY

out = "/Users/thrilzco/Documents/Attenda_Offer_Letter_Drashti.docx"
doc.save(out)
print("saved:", out)