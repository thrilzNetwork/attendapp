#!/usr/bin/env python3
"""Attenda branded offer letter — designed version (logo, gradient rule, callouts)."""
import datetime
from docx import Document
from docx.shared import Pt, Inches, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BRAND = "/tmp/attenda_brand"
TEAL = RGBColor(0x15, 0x8A, 0x7C)
TEAL_HEX = "158A7C"
DARK = RGBColor(0x1A, 0x2A, 0x3A)
GRAY = RGBColor(0x6B, 0x72, 0x80)
MINT_FILL = "E8F4F1"   # soft teal tint for callouts
PALE_FILL = "F4FAF8"

CAND = "Drashti"
TODAY = datetime.date(2026, 9, 2)
ACCEPT_BY = datetime.date(2026, 9, 4)
START = datetime.date(2026, 9, 7)
fmt = lambda d: d.strftime("%B %d, %Y").replace(" 0", " ")

def cell_shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:fill"), fill)
    tcPr.append(shd)

def cell_margins(cell, top=100, bottom=100, left=160, right=160):
    tcPr = cell._tc.get_or_add_tcPr()
    m = OxmlElement("w:tcMar")
    for tag, val in (("top", top), ("bottom", bottom), ("start", left), ("end", right)):
        el = OxmlElement(f"w:{tag}"); el.set(qn("w:w"), str(val)); el.set(qn("w:type"), "dxa")
        m.append(el)
    tcPr.append(m)

def no_table_borders(tbl):
    tblPr = tbl._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}"); el.set(qn("w:val"), "none")
        borders.append(el)
    tblPr.append(borders)

def para_shade(p, fill):
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:fill"), fill)
    pPr.append(shd)

def spacer(pts=8):
    """Controlled gap: short empty line with exact space_after."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(pts)
    pPr = p._p.get_or_add_pPr()
    rPr = OxmlElement("w:rPr")
    sz = OxmlElement("w:sz"); sz.set(qn("w:val"), "12")  # 6pt paragraph mark
    rPr.append(sz); pPr.append(rPr)
    return p

doc = Document()
st = doc.styles["Normal"]
st.font.name = "Calibri"; st.font.size = Pt(10.5); st.font.color.rgb = DARK
st.paragraph_format.space_after = Pt(7)

sec = doc.sections[0]
sec.top_margin = Inches(0.55); sec.bottom_margin = Inches(0.55)
sec.left_margin = Inches(0.75); sec.right_margin = Inches(0.75)

# ============ GRADIENT RULE (top) ============
g = doc.add_paragraph()
g.paragraph_format.space_after = Pt(14)
g.add_run().add_picture(f"{BRAND}/gradient_rule.png", width=Inches(7.0), height=Inches(0.09))

# ============ LOGO ============
lp = doc.add_paragraph()
lp.paragraph_format.space_after = Pt(2)
lp.add_run().add_picture(f"{BRAND}/logo_primary_cropped.png", width=Inches(2.35))

kick = doc.add_paragraph()
kick.paragraph_format.space_after = Pt(10)
rk = kick.add_run("OFFER OF ENGAGEMENT")
rk.font.name = "Georgia"; rk.font.size = Pt(9.5); rk.font.bold = True
rk.font.color.rgb = GRAY
r2 = kick.add_run("     ·     CONFIDENTIAL")
r2.font.name = "Georgia"; r2.font.size = Pt(9.5); r2.font.color.rgb = GRAY
# bottom hairline
pPr = kick._p.get_or_add_pPr()
pBdr = OxmlElement("w:pBdr"); b = OxmlElement("w:bottom")
b.set(qn("w:val"), "single"); b.set(qn("w:sz"), "6"); b.set(qn("w:space"), "6"); b.set(qn("w:color"), "CBE4DF")
pBdr.append(b); pPr.append(pBdr)

# ============ DATE + GREETING ============
d = doc.add_paragraph(fmt(TODAY)); d.paragraph_format.space_after = Pt(12)

h = doc.add_paragraph()
rh = h.add_run(f"Dear {CAND},")
rh.font.name = "Georgia"; rh.font.size = Pt(13); rh.font.bold = True
h.paragraph_format.space_after = Pt(8)

# ============ WELCOME ============
w = doc.add_paragraph(
    "We would like to invite you to join Attenda as our "
)
rw = w.add_run("Controller")
rw.font.bold = True
w.add_run(
    " — a part-time, limited-hours role to start, and the organizational backbone of how this company runs hotels. "
    "You will work directly with me, day to day, keeping every property we serve audited, organized, "
    "and ahead of its calendar."
)

# ============ THE ROLE ============
h = doc.add_paragraph()
h.paragraph_format.space_before = Pt(14); h.paragraph_format.space_after = Pt(6)
rh = h.add_run("What you will own with me")
rh.font.name = "Georgia"; rh.font.size = Pt(14); rh.font.bold = True; rh.font.color.rgb = TEAL

duties = [
    ("Daily audit", "every active property, every morning — occupancy, arrivals and departures, incidents, work orders, and open follow-ups — reviewed and reconciled before the day gets away from anyone."),
    ("Organization", "keeping every property's documents, files, reports, and open items organized, accounted for, and easy to find the moment someone needs them."),
    ("Calendar ownership", "one operational calendar per property — inspections, preventive maintenance, report deadlines, renewals — scheduled, tracked, and never missed."),
    ("Reminders for what matters", "flagging important dates, deadlines, and follow-ups early, so nothing important is discovered after it's already late."),
    ("Monthly owner presentation", "the story of each property — performance, wins, issues, and next month's plan — in a clean format owners actually read."),
]
for title, body in duties:
    bp = doc.add_paragraph(style="List Bullet")
    bp.paragraph_format.space_after = Pt(5)
    rt = bp.add_run(title + " — "); rt.font.bold = True
    bp.add_run(body)

# ============ COMP CALLOUT ============
tbl = doc.add_table(rows=1, cols=1)
no_table_borders(tbl)
cell = tbl.rows[0].cells[0]
cell.width = Inches(7.0)
cell_shade(cell, MINT_FILL)
cell_margins(cell, top=160, bottom=160, left=220, right=220)
cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

cp = cell.paragraphs[0]
cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
cp.paragraph_format.space_after = Pt(2)
r1 = cp.add_run("$600")
r1.font.name = "Georgia"; r1.font.size = Pt(34); r1.font.bold = True; r1.font.color.rgb = TEAL
r2 = cp.add_run("  per hotel, per month")
r2.font.name = "Georgia"; r2.font.size = Pt(13); r2.font.bold = True; r2.font.color.rgb = DARK

cp2 = cell.add_paragraph()
cp2.alignment = WD_ALIGN_PARAGRAPH.CENTER
cp2.paragraph_format.space_after = Pt(4)
r3 = cp2.add_run("A part-time retainer that grows with the portfolio — every hotel Attenda onboards adds $600 to your monthly retainer.")
r3.font.size = Pt(10.5); r3.font.color.rgb = GRAY

cp3 = cell.add_paragraph()
cp3.alignment = WD_ALIGN_PARAGRAPH.CENTER
r4 = cp3.add_run("Attenda is early and growing on purpose — we built this per property so that when we win, you win with us.")
r3b = r4
r4.font.size = Pt(10.5); r4.font.italic = True; r4.font.color.rgb = GRAY

# ============ FACT STRIP ============
spacer(10)
strip = doc.add_table(rows=2, cols=3)
no_table_borders(strip)
facts = [("START DATE", fmt(START)), ("SIGN BY", fmt(ACCEPT_BY)), ("REPORTS TO", "Founder, Attenda")]
for i, (label, value) in enumerate(facts):
    c = strip.rows[0].cells[i]; c2 = strip.rows[1].cells[i]
    for cc in (c, c2):
        cc.width = Inches(2.33)
        cell_shade(cc, PALE_FILL)
        cell_margins(cc, top=80, bottom=80, left=140, right=140)
    rl = c.paragraphs[0].add_run(label)
    rl.font.size = Pt(8.5); rl.font.bold = True; rl.font.color.rgb = GRAY
    rv = c2.paragraphs[0].add_run(value)
    rv.font.size = Pt(11); rv.font.bold = True; rv.font.color.rgb = DARK

# ============ FOUNDER NOTE ============
spacer(10)
q = doc.add_table(rows=1, cols=1)
no_table_borders(q)
qc = q.rows[0].cells[0]
qc.width = Inches(7.0)
cell_shade(qc, PALE_FILL)
cell_margins(qc, top=160, bottom=160, left=220, right=220)
qp = qc.paragraphs[0]
qp.paragraph_format.space_after = Pt(4)
rq = qp.add_run("\u201cYou are not joining a company to watch calendars happen. You are joining to keep the room in order with me \u2014 "
                "hotels only run calmly when somebody has today accounted for and tomorrow already scheduled. That somebody is you.\u201d")
rq.font.name = "Georgia"; rq.font.size = Pt(12.5); rq.font.italic = True; rq.font.color.rgb = DARK
qp2 = qc.add_paragraph()
rq2 = qp2.add_run("— Alejandro, Founder")
rq2.font.size = Pt(10); rq2.font.bold = True; rq2.font.color.rgb = TEAL

# ============ CLOSING + CONFIDENTIALITY ============
doc.add_paragraph(
    "If this sounds like you, sign below and let's get to work. We are excited to build this with you.",
)
conf = doc.add_paragraph(
    "Confidentiality: this engagement includes access to sensitive operational, staffing, and financial information "
    "belonging to Attenda and its client properties. That information is expected to remain strictly confidential, "
    "during and after your engagement."
)
for r in conf.runs: r.font.size = Pt(9); r.font.color.rgb = GRAY
conf.paragraph_format.space_before = Pt(6)

doc.add_paragraph("Warm regards,").paragraph_format.space_after = Pt(14)
sig = doc.add_paragraph()
rs = sig.add_run("Alejandro Soria"); rs.font.bold = True; rs.font.name = "Georgia"; rs.font.size = Pt(13)
doc.add_paragraph("Founder, Attenda").paragraph_format.space_after = Pt(12)

# ============ SIGNATURES ============
tbl2 = doc.add_table(rows=2, cols=2)
no_table_borders(tbl2)
h1, h2 = tbl2.rows[0].cells
b1, b2 = tbl2.rows[1].cells
for cell, label in ((h1, "FOR ATTENDA"), (h2, "ACCEPTANCE")):
    cell.width = Inches(3.5)
    cell_shade(cell, MINT_FILL); cell_margins(cell, top=80, bottom=60, left=160, right=160)
    rr = cell.paragraphs[0].add_run(label)
    rr.font.size = Pt(9); rr.font.bold = True; rr.font.color.rgb = TEAL
for cell, lines in (
    (b1, ["Signature: ____________________", "", "Founder, Attenda", f"Date: {fmt(TODAY)}"]),
    (b2, [f"I, {CAND}, accept this offer.", "", "Signature: ____________________", "", "Date: ______________"]),
):
    cell.width = Inches(3.5); cell_margins(cell, top=100, bottom=100, left=160, right=160)
    first = True
    for text in lines:
        p = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        p.paragraph_format.space_after = Pt(8)
        rr = p.add_run(text); rr.font.size = Pt(10.5)

# ============ FOOTER ============
fp = doc.add_paragraph()
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
fp.paragraph_format.space_before = Pt(12)
fp.add_run().add_picture(f"{BRAND}/gradient_rule.png", width=Inches(7.0), height=Inches(0.09))
fp2 = doc.add_paragraph()
fp2.alignment = WD_ALIGN_PARAGRAPH.CENTER
rf = fp2.add_run("attenda  ·  the operations platform for independent hotels  ·  attendaapp.com")
rf.font.size = Pt(9); rf.font.color.rgb = GRAY

out = "/Users/thrilzco/Documents/Attenda_Offer_Letter_Drashti.docx"
doc.save(out)
print("saved:", out)